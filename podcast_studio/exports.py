"""Readable exports always resolve citations from current project data."""
from __future__ import annotations
import csv
import io
import json
from pathlib import Path
from html import escape
from .model import Project, stamp
from .summaries import evidence

FORMATS = ["TXT", "Markdown", "PDF", "DOCX", "JSON", "CSV", "SRT"]
SCOPES = ["Complete transcript", "Selected speaker transcript", "All summaries", "Selected speaker summaries",
          "Transcript + summaries", "Source-linked notes"]
EXTENSIONS = {"TXT": ".txt", "Markdown": ".md", "PDF": ".pdf", "DOCX": ".docx", "JSON": ".json", "CSV": ".csv", "SRT": ".srt"}


def selection(project, scope, speaker_id=None):
    if scope.startswith("Selected speaker") and speaker_id not in project.speakers:
        raise ValueError("Select a speaker for this export")
    segments = project.segments if scope in (SCOPES[0], SCOPES[4]) else (
        [s for s in project.segments if s.speaker_id == speaker_id] if scope == SCOPES[1] else [])
    summaries = project.summaries if scope in SCOPES[2:] else []
    if scope == SCOPES[3]:
        summaries = [summary for summary in summaries if summary.mode == "Individual speakers"
                     and any(project.segment(i).speaker_id == speaker_id
                             for claim in summary.claims for i in claim.segment_ids
                             if any(s.id == i for s in project.segments))]
    return segments, summaries


def report(project, scope, speaker_id=None, timestamps=True, readable=False):
    segments, summaries = selection(project, scope, speaker_id)
    lines = [project.title, ""]
    if segments:
        lines += ["TRANSCRIPT — " + ("readability copy" if readable else "edited recognized speech"), ""]
    for s in segments:
        when = f"[{stamp(s.start)}–{stamp(s.end)}] " if timestamps else ""
        text = s.readable_text if readable and s.readable_text is not None else s.text
        lines += [f"{when}{project.speakers[s.speaker_id].name}: {text}", ""]
    for summary in summaries:
        lines += [summary.title + (" [STALE — regenerate after corrections]" if summary.stale else ""),
                  summary.provider, ""]
        for claim in summary.claims:
            lines += [f"[{claim.status.upper()}] {claim.text}", claim.note]
            for item in evidence(project, claim):
                lines += [f"Source {item['id']} · {item['speaker']} · {stamp(item['start'])}–{stamp(item['end'])}",
                          "Evidence: " + item["text"]]
            missing = set(claim.segment_ids) - {i["id"] for i in evidence(project, claim)}
            if missing:
                lines.append("Missing references: " + ", ".join(sorted(missing)))
            lines.append("")
    return "\n".join(lines)


def export(project: Project, path: str, format: str, scope: str = SCOPES[0], speaker_id=None,
           timestamps=True, readable=False):
    project.validate()
    segments, summaries = selection(project, scope, speaker_id)
    text = report(project, scope, speaker_id, timestamps, readable)
    destination = Path(path)
    if format in ("TXT", "Markdown"):
        if format == "Markdown":
            # Plain report in a fenced block: preserve literal transcript Markdown and evidence.
            fence = "`" * max(3, max((len(part) for part in text.split("\n") if part and set(part) == {"`"}), default=0) + 1)
            text = f"# {project.title.replace(chr(10), ' ')}\n\n{fence}text\n{text}\n{fence}\n"
        destination.write_text(text, encoding="utf-8")
    elif format == "JSON":
        if scope == SCOPES[0] and not readable:
            project.save(destination)  # Complete reopenable project; other scopes are export views.
        else:
            from dataclasses import asdict
            payload = {"export_schema_version": 1, "project_id": project.id, "title": project.title,
                       "speakers": {k: asdict(v) for k, v in project.speakers.items()},
                       "segments": [asdict(s) for s in segments], "summaries": [asdict(s) for s in summaries]}
            destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    elif format == "CSV":
        with destination.open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.writer(output)
            writer.writerow(["type", "id", "speaker", "start", "end", "text", "status", "sources"])
            def safe(value):
                value = str(value)
                return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value
            for s in segments:
                writer.writerow(["segment", s.id, safe(project.speakers[s.speaker_id].name), stamp(s.start, True),
                                 stamp(s.end, True), safe(s.readable_text if readable and s.readable_text is not None else s.text), "", ""])
            for summary in summaries:
                for claim in summary.claims:
                    writer.writerow(["claim", summary.id, "", "", "", safe(claim.text),
                                     "stale" if summary.stale else claim.status,
                                     json.dumps(evidence(project, claim), ensure_ascii=False)])
    elif format == "SRT":
        if summaries or not segments:
            raise ValueError("SRT is available for transcript scopes; use a report format for summaries")
        blocks = []
        for index, s in enumerate(segments, 1):
            body = s.readable_text if readable and s.readable_text is not None else s.text
            blocks.append(f"{index}\n{stamp(s.start, True).replace('.', ',')} --> {stamp(s.end, True).replace('.', ',')}\n"
                          f"{project.speakers[s.speaker_id].name}: {body}\n")
        destination.write_text("\n".join(blocks), encoding="utf-8")
    elif format == "DOCX":
        from docx import Document
        document = Document()
        document.add_heading(project.title, 0)
        for line in text.splitlines()[1:]:
            document.add_paragraph(line)
        document.save(destination)
    elif format == "PDF":
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        # Bundled ReportLab Vera font handles common Latin accents; CJK/complex scripts need custom fonts.
        import reportlab
        font = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
        pdfmetrics.registerFont(TTFont("StudioVera", str(font)))
        styles = getSampleStyleSheet()
        styles["BodyText"].fontName = "StudioVera"
        styles["BodyText"].wordWrap = "CJK"
        story = []
        for line in text.splitlines():
            story.append(Paragraph(escape(line), styles["BodyText"]) if line else Spacer(1, 8))
        SimpleDocTemplate(str(destination), title=project.title).build(story)
    else:
        raise ValueError("Unknown export format")
